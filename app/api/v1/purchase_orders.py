import csv
import uuid
from datetime import date, datetime, timezone
from io import StringIO

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import get_current_user
from app.models.org import Department
from app.models.purchase_order import PurchaseOrder
from app.models.request import Request
from app.models.user import User
from app.schemas.purchase_order import (
    BatchUploadResponse,
    BatchUploadResult,
    CreatePurchaseOrderPayload,
    PurchaseOrderDetailResponse,
    PurchaseOrderListResponse,
    PurchaseOrderOut,
    UpdatePurchaseOrderPayload,
)
from app.services.po_ledger import committed_amount_for_po, committed_amounts_for_pos

router = APIRouter(prefix="/purchase-orders", tags=["purchase-orders"])

_CURRENCY_LABELS = {
    "naira (ngn)": "NGN",
    "dollar (usd)": "USD",
    "pound (gbp)": "GBP",
    "euro (eur)": "EUR",
    "ngn": "NGN",
    "usd": "USD",
    "gbp": "GBP",
    "eur": "EUR",
}

_JOB_TYPE_LABELS = {
    "work": "work",
    "supply": "supply",
}


def _to_out(
    po: PurchaseOrder, dept_names: dict[str, str], committed: float
) -> PurchaseOrderOut:
    return PurchaseOrderOut(
        id=po.id,
        po_number=po.po_number,
        contractor_name=po.contractor_name,
        contract_amount=po.contract_amount,
        currency=po.currency,
        department_id=po.department_id,
        department_name=dept_names.get(po.department_id),
        description=po.description,
        date_issued=po.date_issued,
        job_type=po.job_type,
        contractor_address=po.contractor_address,
        status=po.status,
        amount_committed=committed,
        amount_remaining=max(po.contract_amount - committed, 0.0),
        created_at=po.created_at,
    )


@router.get("", response_model=PurchaseOrderListResponse)
def list_purchase_orders(
    search: str | None = Query(None),
    status: str | None = Query(None),
    job_type: str | None = Query(None),
    department_id: str | None = Query(None),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
    db: Session = Depends(get_db),
    _current_user: User = Depends(get_current_user),
):
    query = db.query(PurchaseOrder)
    if status and status != "all":
        query = query.filter(PurchaseOrder.status == status)
    if job_type:
        # Historical data has inconsistent casing (Work/work/Service) — compare
        # case-insensitively. POs with no value set are treated as matching any
        # filter instead of being hidden outright.
        query = query.filter(
            (PurchaseOrder.job_type.ilike(job_type))
            | (PurchaseOrder.job_type.is_(None))
        )
    if department_id:
        query = query.filter(PurchaseOrder.department_id == department_id)
    if search:
        like = f"%{search.strip()}%"
        matching_dept_ids = [
            d.id for d in db.query(Department).filter(Department.name.ilike(like)).all()
        ]
        query = query.filter(
            (PurchaseOrder.po_number.ilike(like))
            | (PurchaseOrder.contractor_name.ilike(like))
            | (PurchaseOrder.department_id.in_(matching_dept_ids))
        )

    total = query.count()
    purchase_orders = (
        # Most recently issued first; POs without a date fall to the end.
        query.order_by(
            PurchaseOrder.date_issued.is_(None), PurchaseOrder.date_issued.desc()
        )
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )
    dept_names = {d.id: d.name for d in db.query(Department).all()}
    committed_by_po = committed_amounts_for_pos(db, [po.id for po in purchase_orders])
    items = [_to_out(po, dept_names, committed_by_po[po.id]) for po in purchase_orders]
    return PurchaseOrderListResponse(
        data=items, meta={"total": total, "page": page, "pageSize": page_size}
    )


@router.post("", status_code=201, response_model=PurchaseOrderDetailResponse)
def create_purchase_order(
    payload: CreatePurchaseOrderPayload,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if (
        db.query(PurchaseOrder)
        .filter(PurchaseOrder.po_number == payload.po_number)
        .first()
    ):
        raise HTTPException(
            status_code=409, detail="A PO with this number already exists."
        )

    department = db.get(Department, payload.department_id)
    if department is None:
        raise HTTPException(status_code=400, detail="Unknown department.")

    po = PurchaseOrder(
        id=uuid.uuid4().hex,
        po_number=payload.po_number.strip(),
        contractor_name=payload.contractor_name.strip(),
        contract_amount=payload.contract_amount,
        currency=payload.currency,
        department_id=payload.department_id,
        description=payload.description,
        date_issued=payload.date_issued,
        job_type=payload.job_type,
        contractor_address=payload.contractor_address,
        status=payload.status,
        created_by_id=current_user.id,
        created_at=datetime.now(timezone.utc),
    )
    db.add(po)
    db.commit()
    db.refresh(po)

    dept_names = {department.id: department.name}
    committed = committed_amount_for_po(db, po.id)
    return PurchaseOrderDetailResponse(data=_to_out(po, dept_names, committed))


@router.get("/export/csv")
def export_purchase_orders_csv(
    search: str | None = Query(None),
    status: str | None = Query(None),
    db: Session = Depends(get_db),
    _current_user: User = Depends(get_current_user),
):
    query = db.query(PurchaseOrder)
    if status and status != "all":
        query = query.filter(PurchaseOrder.status == status)
    if search:
        like = f"%{search.strip()}%"
        query = query.filter(
            (PurchaseOrder.po_number.ilike(like))
            | (PurchaseOrder.contractor_name.ilike(like))
        )
    purchase_orders = query.order_by(PurchaseOrder.created_at.desc()).all()
    dept_names = {d.id: d.name for d in db.query(Department).all()}
    committed_by_po = committed_amounts_for_pos(db, [po.id for po in purchase_orders])

    buf = StringIO()
    writer = csv.writer(buf)
    writer.writerow(
        [
            "PO Number",
            "Contractor Name",
            "Contractor Address",
            "Contract Amount",
            "Currency",
            "Department",
            "Job Type",
            "Description",
            "Date Issued",
            "Status",
            "Amount Committed",
            "Amount Remaining",
            "Created At",
        ]
    )
    for po in purchase_orders:
        committed = committed_by_po[po.id]
        writer.writerow(
            [
                po.po_number,
                po.contractor_name,
                po.contractor_address or "",
                po.contract_amount,
                po.currency,
                dept_names.get(po.department_id, po.department_id),
                po.job_type or "",
                po.description or "",
                po.date_issued or "",
                po.status,
                committed,
                max(po.contract_amount - committed, 0.0),
                po.created_at.isoformat(),
            ]
        )

    filename = (
        f"Kora-purchase-orders-{datetime.now(timezone.utc).date().isoformat()}.csv"
    )
    return Response(
        content=buf.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _parse_currency(raw: str | None) -> str | None:
    if not raw:
        return None
    return _CURRENCY_LABELS.get(raw.strip().lower(), raw.strip().upper())


def _parse_job_type(raw: str | None) -> str | None:
    if not raw:
        return None
    return _JOB_TYPE_LABELS.get(raw.strip().lower())


def _parse_cell_date(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    return str(value).strip()


@router.post("/batch-upload", response_model=BatchUploadResponse)
def batch_upload_purchase_orders(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=400, detail="Please upload the .xlsx PO upload template."
        )

    try:
        workbook = openpyxl.load_workbook(file.file, data_only=True)
    except Exception as exc:
        raise HTTPException(
            status_code=400, detail=f"Could not read the file: {exc}"
        ) from exc

    sheet = (
        workbook["PO Upload Template"]
        if "PO Upload Template" in workbook.sheetnames
        else workbook.worksheets[0]
    )
    departments_by_name = {
        d.name.strip().lower(): d for d in db.query(Department).all()
    }

    created = 0
    skipped: list[str] = []
    now = datetime.now(timezone.utc)

    # Row 1 = headers, row 2 = filled-in example — real data starts row 3.
    for row_index, row in enumerate(
        sheet.iter_rows(min_row=3, values_only=True), start=3
    ):
        (
            po_number,
            contractor_name,
            contract_amount,
            currency_raw,
            department_name,
            description,
            date_issued_raw,
            status_raw,
            job_type_raw,
            contractor_address,
        ) = (list(row) + [None] * 10)[:10]

        if not po_number and not contractor_name:
            continue  # blank row

        po_number = str(po_number).strip() if po_number else ""
        if not po_number:
            skipped.append(f"Row {row_index}: missing PO number.")
            continue
        if not contractor_name:
            skipped.append(f"Row {row_index} ({po_number}): missing contractor name.")
            continue
        if not contract_amount:
            skipped.append(f"Row {row_index} ({po_number}): missing contract amount.")
            continue
        if db.query(PurchaseOrder).filter(PurchaseOrder.po_number == po_number).first():
            skipped.append(
                f"Row {row_index} ({po_number}): a PO with this number already exists."
            )
            continue

        department = departments_by_name.get(str(department_name or "").strip().lower())
        if department is None:
            skipped.append(
                f"Row {row_index} ({po_number}): unknown department '{department_name}'."
            )
            continue

        currency = _parse_currency(str(currency_raw) if currency_raw else None) or "NGN"
        status_value = str(status_raw).strip().lower() if status_raw else "active"
        if status_value not in ("active", "closed"):
            status_value = "active"

        po = PurchaseOrder(
            id=uuid.uuid4().hex,
            po_number=po_number,
            contractor_name=str(contractor_name).strip(),
            contract_amount=float(contract_amount),
            currency=currency,
            department_id=department.id,
            description=str(description).strip() if description else None,
            date_issued=_parse_cell_date(date_issued_raw),
            status=status_value,
            job_type=_parse_job_type(str(job_type_raw) if job_type_raw else None),
            contractor_address=str(contractor_address).strip()
            if contractor_address
            else None,
            created_by_id=current_user.id,
            created_at=now,
        )
        db.add(po)
        created += 1

    db.commit()
    return BatchUploadResponse(data=BatchUploadResult(created=created, skipped=skipped))


@router.get("/{po_id}", response_model=PurchaseOrderDetailResponse)
def get_purchase_order(
    po_id: str,
    db: Session = Depends(get_db),
    _current_user: User = Depends(get_current_user),
):
    po = db.get(PurchaseOrder, po_id)
    if po is None:
        raise HTTPException(status_code=404, detail="Purchase order not found.")
    dept_names = {d.id: d.name for d in db.query(Department).all()}
    committed = committed_amount_for_po(db, po.id)
    return PurchaseOrderDetailResponse(data=_to_out(po, dept_names, committed))


@router.patch("/{po_id}", response_model=PurchaseOrderDetailResponse)
def update_purchase_order(
    po_id: str,
    payload: UpdatePurchaseOrderPayload,
    db: Session = Depends(get_db),
    _current_user: User = Depends(get_current_user),
):
    po = db.get(PurchaseOrder, po_id)
    if po is None:
        raise HTTPException(status_code=404, detail="Purchase order not found.")

    if payload.po_number and payload.po_number != po.po_number:
        if (
            db.query(PurchaseOrder)
            .filter(
                PurchaseOrder.po_number == payload.po_number, PurchaseOrder.id != po_id
            )
            .first()
        ):
            raise HTTPException(
                status_code=409, detail="A PO with this number already exists."
            )
        po.po_number = payload.po_number.strip()

    if payload.contractor_name is not None:
        po.contractor_name = payload.contractor_name.strip()
    if payload.contract_amount is not None:
        po.contract_amount = payload.contract_amount
    if payload.currency is not None:
        po.currency = payload.currency
    if payload.department_id is not None:
        department = db.get(Department, payload.department_id)
        if department is None:
            raise HTTPException(status_code=400, detail="Unknown department.")
        po.department_id = payload.department_id
    if payload.description is not None:
        po.description = payload.description
    if payload.date_issued is not None:
        po.date_issued = payload.date_issued
    if payload.job_type is not None:
        po.job_type = payload.job_type
    if payload.contractor_address is not None:
        po.contractor_address = payload.contractor_address
    if payload.status is not None:
        po.status = payload.status

    db.commit()
    db.refresh(po)
    dept_names = {d.id: d.name for d in db.query(Department).all()}
    committed = committed_amount_for_po(db, po.id)
    return PurchaseOrderDetailResponse(data=_to_out(po, dept_names, committed))


@router.delete("/{po_id}", status_code=204)
def delete_purchase_order(
    po_id: str,
    db: Session = Depends(get_db),
    _current_user: User = Depends(get_current_user),
):
    po = db.get(PurchaseOrder, po_id)
    if po is None:
        raise HTTPException(status_code=404, detail="Purchase order not found.")

    if db.query(Request).filter(Request.po_id == po_id).first():
        raise HTTPException(
            status_code=409,
            detail="This PO has requests submitted against it and can't be deleted — mark it Closed instead.",
        )

    db.delete(po)
    db.commit()
